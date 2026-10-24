# -*- coding: utf-8 -*-
# 原始件生成脚本：确定性生成三类演示原始件至 data/raw/
#   - D-300操作手册.pdf        ：reportlab 生成，3 页，第 2 页含带框线参数表格，
#                                第 1 页含装饰小图（应被过滤），第 3 页含带 caption 的示意图（应被抽取）
#   - E-450保养手册.docx       ：python-docx 生成，含保养表格与显式分页符
#   - 维修记录单_scan/*.png    ：PIL 渲染中文文本，加高斯噪声与微倾斜模拟扫描件
# PDF 插图用 PIL 预生成 PNG 至 data/raw/_assets/（get_image_info 只识别位图，
# reportlab 矢量 Drawing 无法被识别，故插图必须先落位图再嵌入）
# 所有随机源均使用固定种子，重复运行产物内容一致（确定性生成）
from pathlib import Path

import numpy as np

# 项目根目录（scripts/ 的上一级）
ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "data" / "raw"
ASSETS_DIR = RAW_DIR / "_assets"

# Windows 系统中文字体（微软雅黑 TTC）
FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"


def _gen_pdf_assets() -> tuple[Path, Path]:
    # 生成 PDF 插图位图资产（PIL 绘制，无随机源，字节级确定）：
    #   大图：液压系统动力单元示意图（白底，方框 + 连线 + 中文标注）
    #   小图：装饰用实心圆角标（应被面积阈值过滤）
    from PIL import Image, ImageDraw, ImageFont

    ASSETS_DIR.mkdir(parents=True, exist_ok=True)

    # 示意图：800x500 白底，油箱 -> 油泵 -> 液压缸 三框两线
    diagram = Image.new("RGB", (800, 500), (255, 255, 255))
    draw = ImageDraw.Draw(diagram)
    font = ImageFont.truetype(FONT_PATH, 36, index=0)
    boxes = [((60, 200), (240, 320), "油箱"),
             ((340, 200), (520, 320), "油泵"),
             ((620, 200), (780, 320), "液压缸")]
    for (x0, y0), (x1, y1), label in boxes:
        draw.rectangle([x0, y0, x1, y1], outline=(0, 0, 0), width=4)
        # 标注文字大致居中于框内
        tw = draw.textlength(label, font=font)
        draw.text(((x0 + x1 - tw) / 2, (y0 + y1) / 2 - 22), label,
                  font=font, fill=(0, 0, 0))
    # 框间连线（带简单箭头）
    for x_start, x_end in ((240, 340), (520, 620)):
        draw.line([x_start, 260, x_end, 260], fill=(0, 0, 0), width=4)
        draw.polygon([(x_end, 260), (x_end - 16, 250), (x_end - 16, 270)],
                     fill=(0, 0, 0))
    diagram_path = ASSETS_DIR / "hydraulic_diagram.png"
    diagram.save(diagram_path)

    # 装饰小图：40x40 实心圆（无文字）
    decor = Image.new("RGB", (40, 40), (255, 255, 255))
    ImageDraw.Draw(decor).ellipse([4, 4, 36, 36], fill=(90, 90, 90))
    decor_path = ASSETS_DIR / "decor_dot.png"
    decor.save(decor_path)

    return diagram_path, decor_path


def gen_pdf() -> Path:
    # 生成 D-300 操作手册 PDF：3 页，第 2 页含带完整框线的参数表格，
    # 第 1 页含装饰小图（< 页面 4%），第 3 页含带 caption 的示意图大图（>= 页面 4%）
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image,
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors

    # 确定性输出：invariant=1 时 CreationDate 与文档 ID 固定，字节级可复现
    import reportlab.rl_config as rl_config
    rl_config.invariant = 1

    # 插图位图资产（每次重新生成，PIL 绘制无随机源，保证确定性）
    diagram_path, decor_path = _gen_pdf_assets()

    # 注册 TTC 中文字体需指定 subfontIndex=0
    pdfmetrics.registerFont(TTFont("msyh", FONT_PATH, subfontIndex=0))
    style = ParagraphStyle("cn", fontName="msyh", fontSize=12, leading=20)
    title_style = ParagraphStyle("cn_title", fontName="msyh", fontSize=18, leading=28)
    caption_style = ParagraphStyle("cn_caption", fontName="msyh", fontSize=10, leading=16)

    out = RAW_DIR / "D-300操作手册.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            topMargin=25 * mm, bottomMargin=25 * mm)
    story = []

    # 第 1 页：概述（末尾放 40x40pt 装饰小图，面积约 0.3% < 4%，应被过滤）
    story.append(Paragraph("D-300 注塑机操作手册（V1.0）", title_style))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("本手册适用于 D-300 型注塑机的操作与日常点检。", style))
    story.append(Paragraph("设备液压系统为高压系统，作业前必须确认系统压力已卸放。", style))
    story.append(Paragraph("开机前检查油箱油位，油位应处于油标上下限之间。", style))
    story.append(Spacer(1, 10 * mm))
    story.append(Image(str(decor_path), width=40, height=40))
    story.append(PageBreak())

    # 第 2 页：技术参数表格（完整框线，便于 find_tables 识别）
    story.append(Paragraph("技术参数表", title_style))
    story.append(Spacer(1, 6 * mm))
    table_data = [
        ["参数名称", "数值", "单位", "备注"],
        ["锁模力", "3000", "kN", "四拉杆直压式锁模"],
        # 表格内容避免全角标点、波浪号、小数点：reportlab 绘制顺序特殊，PyMuPDF 抽取会错位
        ["液压系统压力上限", "18", "MPa", "安全相关参数 严禁超调"],
        ["液压系统额定工作压力", "15", "MPa", "正常生产设定范围 13至15"],
        ["射胶量（PS 料）", "520", "g", "理论射胶容积 580 cm3"],
        ["螺杆直径", "55", "mm", "标配 B 组螺杆"],
        ["油箱容量", "300", "L", "推荐 ISO VG 46 抗磨液压油"],
        ["整机重量", "8600", "kg", "不含液压油与模具"],
    ]
    tbl = Table(table_data, colWidths=[55 * mm, 25 * mm, 20 * mm, 55 * mm])
    tbl.setStyle(TableStyle([
        # 完整框线：find_tables 依赖连续线段识别表格
        ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "msyh"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("ALIGN", (1, 0), (2, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(tbl)
    story.append(PageBreak())

    # 第 3 页：操作注意事项 + 示意图大图（400x250pt，面积约 20% >= 4%，应被抽取）
    # caption 紧跟图下方（无 Spacer），保证与图底边 y 距离在阈值内
    story.append(Paragraph("操作注意事项", title_style))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("1. 启动油泵电机后空载运行 3~5 分钟，空载压力应低于 2MPa。", style))
    story.append(Paragraph("2. 生产状态下系统压力不得超过 18MPa 上限。", style))
    story.append(Paragraph("3. 液压油温度正常工作范围为 30~55 摄氏度，超过 60 摄氏度应停机检查。", style))
    story.append(Paragraph("4. 更换模具时必须按下急停按钮并挂警示牌。", style))
    story.append(Spacer(1, 6 * mm))
    story.append(Image(str(diagram_path), width=400, height=250))
    story.append(Paragraph("图 1-1 液压系统动力单元示意图", caption_style))

    doc.build(story)
    return out


def gen_docx() -> Path:
    # 生成 E-450 保养手册 DOCX：2 个显式分页符共 3 页，含保养项目表格
    from docx import Document
    from docx.enum.text import WD_BREAK

    out = RAW_DIR / "E-450保养手册.docx"
    document = Document()

    # 第 1 页：概述（含可断言数值）
    document.add_heading("E-450 注塑机保养手册（V1.0）", level=1)
    document.add_paragraph("本手册适用于 E-450 型注塑机的定期保养作业。")
    document.add_paragraph("液压油更换周期 2000 小时，更换时同步清理油箱底部沉积物。")
    document.add_paragraph("首次保养应在设备运行 500 小时后进行。")
    # 显式分页符
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 第 2 页：保养项目表格
    document.add_heading("定期保养项目表", level=1)
    table = document.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    rows = [
        ["保养项目", "周期", "工时", "备注"],
        ["液压油更换", "2000 小时", "2.0 h", "使用 ISO VG 46 抗磨液压油"],
        ["回油滤芯更换", "1000 小时", "0.5 h", "更换时检查滤芯压差"],
        ["螺杆料筒检查", "4000 小时", "4.0 h", "测量螺杆与料筒间隙"],
        ["锁模机构润滑", "500 小时", "0.5 h", "加注锂基润滑脂"],
        ["安全门联锁测试", "250 小时", "0.2 h", "安全相关项目，不得遗漏"],
    ]
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            table.rows[i].cells[j].text = text
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 第 3 页：注意事项
    document.add_heading("保养注意事项", level=1)
    document.add_paragraph("1. 保养作业前必须停机卸压，确认系统压力表指针回落至零。")
    document.add_paragraph("2. 液压油加油量为 280 升，油位应处于油标上下限之间。")
    document.add_paragraph("3. 保养完成后填写保养记录，并由班组长签字确认。")

    document.save(str(out))
    _fix_zip_timestamps(out)
    return out


def _fix_zip_timestamps(path: Path) -> None:
    # docx 内部 XML 内容本身确定，仅 zip 条目时间戳随打包时间变化；
    # 重打包并固定条目时间，保证字节级可复现
    import zipfile
    fixed_dt = (2026, 1, 1, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as zin:
        items = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=fixed_dt)
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


# 扫描件维修记录文本（一页一条，纯文本无表格）
SCAN_RECORDS = [
    [
        "维修记录单",
        "日期：2026-03-05",
        "设备型号：A-160",
        "故障现象：射胶压力波动，制品重量偏差超过 3 克。",
        "处理结果：更换射胶油缸密封件，试机后压力恢复稳定。",
        "维修人员：张工",
    ],
    [
        "维修记录单",
        "日期：2026-04-12",
        "设备型号：A-160",
        "故障现象：液压油温升至 68 摄氏度并触发高温报警。",
        "处理结果：清理冷却器水路水垢，油温回落至 45 摄氏度。",
        "维修人员：李工",
    ],
    [
        "维修记录单",
        "日期：2026-05-20",
        "设备型号：A-160",
        "故障现象：锁模动作异响，锁模力无法达到设定值。",
        "处理结果：紧固哥林柱螺母并补充润滑脂，锁模力恢复正常。",
        "维修人员：王工",
    ],
]


def gen_scan_images() -> list[Path]:
    # 生成维修记录单扫描件：PIL 渲染中文文本为 PNG，加高斯噪声与微倾斜
    from PIL import Image, ImageDraw, ImageFont

    out_dir = RAW_DIR / "维修记录单_scan"
    out_dir.mkdir(parents=True, exist_ok=True)

    font_title = ImageFont.truetype(FONT_PATH, 40, index=0)
    font_body = ImageFont.truetype(FONT_PATH, 30, index=0)
    # 固定倾角（0.5 度以内），保证确定性
    rotate_angle = 0.3
    # 固定噪声强度与随机种子
    noise_sigma = 6.0

    paths = []
    for idx, lines in enumerate(SCAN_RECORDS, start=1):
        width, height = 1000, 620
        img = Image.new("L", (width, height), 255)
        draw = ImageDraw.Draw(img)
        y = 60
        for k, line in enumerate(lines):
            font = font_title if k == 0 else font_body
            draw.text((80, y), line, font=font, fill=0)
            y += 70 if k == 0 else 78

        # 整体微倾斜：旋转时以白色填充空白
        img = img.rotate(rotate_angle, resample=Image.BICUBIC, fillcolor=255)

        # 加少量高斯噪声：固定种子保证重复生成一致
        rng = np.random.default_rng(seed=20261024 + idx)
        arr = np.asarray(img).astype(np.float32)
        arr += rng.normal(0.0, noise_sigma, size=arr.shape).astype(np.float32)
        img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8), mode="L")

        out = out_dir / f"page_{idx:03d}.png"
        img.save(out)
        paths.append(out)
    return paths


def main() -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    pdf = gen_pdf()
    print(f"[生成] {pdf}")
    docx = gen_docx()
    print(f"[生成] {docx}")
    imgs = gen_scan_images()
    for p in imgs:
        print(f"[生成] {p}")
    print("全部原始件生成完成")


if __name__ == "__main__":
    main()
